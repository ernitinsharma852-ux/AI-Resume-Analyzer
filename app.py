
import os
import re
import hashlib
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape

import streamlit as st
import pymupdf
from docx import Document
from dotenv import load_dotenv
from google import genai

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)

# -------------------- Configuration --------------------

BASE_DIR = Path(__file__).resolve().parent
load_dotenv(BASE_DIR / ".env")

MAX_FILE_SIZE = 20 * 1024 * 1024
GEMINI_MODEL = os.getenv("GEMINI_MODEL", "gemini-3.8-flash")

st.set_page_config(
    page_title="AI Resume Analyzer",
    page_icon="📄",
    layout="wide",
)

# Load external CSS
css_path = BASE_DIR / "style.css"
if css_path.exists():
    css = css_path.read_text(encoding="utf-8")
    st.markdown(f"<style>{css}</style>", unsafe_allow_html=True)

# -------------------- Session State --------------------

for key, default in {
    "analysis": None,
    "ai_feedback": None,
    "analysis_signature": None,
}.items():
    if key not in st.session_state:
        st.session_state[key] = default

# -------------------- Skills --------------------

SKILLS = [
    "python", "java", "c++", "c", "sql", "mysql",
    "machine learning", "deep learning", "artificial intelligence",
    "data analysis", "data science", "pandas", "numpy",
    "scikit-learn", "tensorflow", "pytorch", "nlp",
    "excel", "power bi", "tableau", "git", "github",
    "rest api", "api", "html", "css", "javascript",
    "react", "node.js", "flask", "django", "streamlit",
    "aws", "azure", "docker", "linux", "communication",
    "problem solving", "iot", "robotics", "opencv",
]

# -------------------- Text Extraction --------------------

def extract_resume_text(uploaded_file):
    data = uploaded_file.getvalue()
    filename = uploaded_file.name.lower()

    if filename.endswith(".pdf"):
        with pymupdf.open(stream=data, filetype="pdf") as pdf:
            text = "\n".join(page.get_text() for page in pdf)

    elif filename.endswith(".docx"):
        doc = Document(BytesIO(data))
        parts = [p.text for p in doc.paragraphs]

        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(cell.text for cell in row.cells))

        text = "\n".join(parts)

    else:
        raise ValueError("Please upload a PDF or DOCX file.")

    if not text.strip():
        raise ValueError(
            "No readable text found. If this is a scanned PDF, "
            "OCR may be required."
        )

    return text


def find_skills(text):
    text_lower = text.lower()
    found = []

    for skill in SKILLS:
        pattern = r"(?<!\w)" + re.escape(skill) + r"(?!\w)"
        if re.search(pattern, text_lower):
            found.append(skill)

    return sorted(set(found))


def calculate_score(resume_text, job_description):
    resume_skills = find_skills(resume_text)
    job_skills = find_skills(job_description)

    if not job_description.strip():
        return {
            "score": 0,
            "resume_skills": resume_skills,
            "matched_skills": [],
            "missing_skills": [],
            "job_skills": [],
        }

    if not job_skills:
        return {
            "score": 0,
            "resume_skills": resume_skills,
            "matched_skills": [],
            "missing_skills": [],
            "job_skills": [],
        }

    matched = sorted(set(resume_skills) & set(job_skills))
    missing = sorted(set(job_skills) - set(resume_skills))
    score = round(len(matched) / len(job_skills) * 100)

    return {
        "score": score,
        "resume_skills": resume_skills,
        "matched_skills": matched,
        "missing_skills": missing,
        "job_skills": job_skills,
    }

# -------------------- Gemini AI Feedback --------------------

def generate_ai_feedback(resume_text, job_description):
    api_key = os.getenv("GEMINI_API_KEY")

    if not api_key:
        raise ValueError(
            "GEMINI_API_KEY not found. Add it to your .env file."
        )

    prompt = f"""
You are a professional resume reviewer and career coach.

Analyze the resume against the job description.

Provide:
1. Overall resume assessment.
2. Strong points.
3. Missing skills or experience.
4. Specific improvement suggestions.
5. ATS-friendly keywords to consider.
6. Suggestions for improving resume bullet points.
7. A short action plan.

Be constructive and accurate. Do not invent qualifications or experience.
Clearly state that this is guidance, not a hiring decision.

RESUME:
{resume_text[:18000]}

JOB DESCRIPTION:
{job_description[:8000]}
"""

    client = genai.Client(api_key=api_key)
    try:
        response = client.models.generate_content(
            model=GEMINI_MODEL,
            contents=prompt,
        )
        if not response.text:
            raise ValueError("The AI returned an empty response.")
        return response.text
    finally:
        client.close()

# -------------------- PDF Report --------------------

def create_pdf_report(analysis, feedback, filename):
    buffer = BytesIO()

    pdf = SimpleDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=42,
        leftMargin=42,
        topMargin=42,
        bottomMargin=42,
        title="AI Resume Analyzer Report",
    )

    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(
        name="ReportTitle",
        parent=styles["Title"],
        alignment=TA_CENTER,
        textColor=colors.HexColor("#4169a1"),
        spaceAfter=18,
    ))
    styles.add(ParagraphStyle(
        name="SectionHeading",
        parent=styles["Heading2"],
        textColor=colors.HexColor("#4169a1"),
        spaceBefore=12,
        spaceAfter=7,
    ))

    story = [
        Paragraph("AI Resume Analyzer Report", styles["ReportTitle"]),
        Paragraph(f"Resume: {escape(filename)}", styles["Normal"]),
        Spacer(1, 12),
        Paragraph("Resume Match Score", styles["SectionHeading"]),
        Paragraph(
            f"<b>{analysis['score']}%</b>",
            styles["Heading1"],
        ),
        Paragraph("Detected Resume Skills", styles["SectionHeading"]),
        Paragraph(
            escape(", ".join(analysis["resume_skills"]) or "None detected"),
            styles["BodyText"],
        ),
        Paragraph("Matched Skills", styles["SectionHeading"]),
        Paragraph(
            escape(", ".join(analysis["matched_skills"]) or "None detected"),
            styles["BodyText"],
        ),
        Paragraph("Missing Skills", styles["SectionHeading"]),
        Paragraph(
            escape(", ".join(analysis["missing_skills"]) or "None detected"),
            styles["BodyText"],
        ),
    ]

    if feedback:
        story.append(Paragraph("AI Feedback", styles["SectionHeading"]))

        for line in feedback.splitlines():
            line = line.strip()
            if line:
                story.append(
                    Paragraph(escape(line), styles["BodyText"])
                )
                story.append(Spacer(1, 5))

    story.extend([
        Spacer(1, 18),
        Paragraph(
            "This report provides automated guidance only. "
            "It is not a hiring decision or a guarantee of selection.",
            styles["Italic"],
        ),
    ])

    pdf.build(story)
    buffer.seek(0)
    return buffer.getvalue()

# -------------------- User Interface --------------------

st.title("📄 AI Resume Analyzer")
st.caption(
    "Upload your resume, compare skills with a job description, "
    "and receive AI-powered suggestions."
)

st.info(
    "Supported formats: PDF and DOCX. Maximum file size: 20 MB."
)

uploaded_file = st.file_uploader(
    "Upload your resume",
    type=["pdf", "docx"],
    max_upload_size=20,
)

if uploaded_file:
    if uploaded_file.size > MAX_FILE_SIZE:
        st.error("File is larger than 20 MB. Please upload a smaller file.")
        uploaded_file = None
    else:
        st.success(
            f"Selected: {uploaded_file.name} "
            f"({uploaded_file.size / (1024 * 1024):.2f} MB)"
        )

st.subheader("💼 Job Description")
job_description = st.text_area(
    "Paste the job description here",
    placeholder=(
        "Example: Python Developer Intern with knowledge of "
        "Python, SQL, Machine Learning, Git, and REST APIs."
    ),
    height=160,
)

if st.button(
    "Analyze Resume",
    type="primary",
    use_container_width=True,
    key="analyze_resume"
):    
    if not uploaded_file:
        st.error("Please upload a PDF or DOCX resume first.")
    elif not job_description.strip():
        st.error("Please paste the job description first.")
    else:
        try:
            with st.spinner("Analyzing your resume..."):
                resume_text = extract_resume_text(uploaded_file)
                signature = hashlib.sha256(
                    (
                        uploaded_file.name
                        + resume_text
                        + job_description
                    ).encode("utf-8")
                ).hexdigest()

                if signature != st.session_state.analysis_signature:
                    st.session_state.ai_feedback = None

                st.session_state.analysis = calculate_score(
                    resume_text, job_description
                )
                st.session_state.analysis["resume_text"] = resume_text
                st.session_state.analysis["job_description"] = job_description
                st.session_state.analysis["filename"] = uploaded_file.name
                st.session_state.analysis_signature = signature

            st.success("Resume analysis completed!")

        except Exception as exc:
            st.error(f"Could not analyze the resume: {exc}")

analysis = st.session_state.analysis

if analysis:
    st.divider()
    st.subheader("📊 Resume Analysis")

    col1, col2, col3 = st.columns(3)
    col1.metric("Skill Match Score", f"{analysis['score']}%")
    col2.metric("Matched Skills", len(analysis["matched_skills"]))
    col3.metric("Missing Skills", len(analysis["missing_skills"]))

    st.subheader("✅ Matched Skills")
    if analysis["matched_skills"]:
        st.write(", ".join(analysis["matched_skills"]))
    else:
        st.write("No matching skills detected from the supported list.")

    st.subheader("📌 Missing Skills")
    if analysis["missing_skills"]:
        st.write(", ".join(analysis["missing_skills"]))
    else:
        st.write("No missing skills detected from the supported list.")

    st.subheader("🧠 AI Resume Feedback")
    st.caption(
        "When you generate AI feedback, extracted resume text and "
        "the job description are sent to Google's Gemini API."
    )

    if st.button("✨ Generate AI Feedback", use_container_width=True):
        try:
            with st.spinner("Generating AI feedback..."):
                st.session_state.ai_feedback = generate_ai_feedback(
                    analysis["resume_text"],
                    analysis["job_description"],
                )
        except Exception as exc:
            st.error(
                "Gemini could not generate feedback. Check the API key, "
                "model availability, internet connection, and API quota."
            )
            st.caption(f"Details: {exc}")

    if st.session_state.ai_feedback:
        st.markdown(st.session_state.ai_feedback)

    report_bytes = create_pdf_report(
        analysis,
        st.session_state.ai_feedback,
        analysis["filename"],
    )

    st.download_button(
        "⬇️ Download PDF Report",
        data=report_bytes,
        file_name="AI_Resume_Analysis_Report.pdf",
        mime="application/pdf",
        use_container_width=True,
    )

# -------------------- Footer --------------------

st.markdown(
    """
    <div class="footer">
        <div class="footer-line"></div>
        <p>© 2026 AI Resume Analyzer</p>
        <p>Developed by Nitin Sharma</p>
    </div>
    """,
    unsafe_allow_html=True,
)
